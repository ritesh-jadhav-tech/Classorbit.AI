from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.parsers.base import ParsedCell, ParsedDocument
from app.services.import_service import ImportService


def test_multi_sheet_excel_creates_separate_class_timetables(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "FYBA A"
    second = workbook.create_sheet("SYBA B")
    for sheet, class_name, division, subject, room in (
        (first, "FYBA", "A", "English", "101"),
        (second, "SYBA", "B", "History", "102"),
    ):
        sheet.append([f"Class: {class_name}", f"Div: {division}", ""])
        sheet.append(["Time", "Monday", "Tuesday"])
        sheet.append(["8:15 AM - 9:15 AM", f"{subject} {room}", "-"])
    output = BytesIO()
    workbook.save(output)

    response = client.post(
        "/api/imports",
        headers=admin_headers,
        files={
            "file": (
                "college-timetables.xlsx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201, response.text
    assert response.json()["summary_json"]["generated_timetable_entries"] == 2

    structures = client.get("/api/timetable-structures", headers=admin_headers).json()
    contexts = sorted(
        (item["timetable"]["class_name"], item["timetable"]["division"], len(item["entries"]))
        for item in structures
    )
    assert contexts == [("FYBA", "A", 1), ("SYBA", "B", 1)]


def test_marker_based_excel_grid_inherits_context_and_publishes_authoritative_overlaps(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in (
        ["CLASSORBIT TIMETABLE TEMPLATE", "", ""],
        ["Department", "Performing Arts", ""],
        ["Class / Program", "SYBA", ""],
        ["Academic Year", "2026-2027", ""],
        ["Time", "Monday", "Tuesday"],
        ["8:15 AM - 9:15 AM", "History 102", ""],
        ["8:15 AM - 9:15 AM", "Performing Arts 102", ""],
    ):
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)

    response = client.post(
        "/api/imports",
        headers=admin_headers,
        files={
            "file": (
                "legacy-classorbit-grid.xlsx",
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201, response.text
    summary = response.json()["summary_json"]
    assert response.json()["status"] == "COMPLETED"
    assert summary["valid"] == 2
    assert summary["needs_review"] == 0
    assert summary["conflicts"] > 0
    assert len(client.get("/api/timetable-entries", headers=admin_headers).json()) == 2


def test_csv_import_maps_room_and_answers_schedule(client: TestClient, admin_headers: dict[str, str]) -> None:
    classroom_response = client.post(
        "/api/classrooms",
        headers=admin_headers,
        json={"room_number": "414", "display_name": "Lecture Hall 414", "capacity": 80},
    )
    assert classroom_response.status_code == 201, classroom_response.text

    csv_content = (
        "class_name,division,day,start_time,end_time,subject,room,batch\n"
        "FYBA,A,Monday,07:15,08:15,English,414,B1\n"
    ).encode()
    import_response = client.post(
        "/api/imports",
        headers=admin_headers,
        files={"file": ("fyba.csv", csv_content, "text/csv")},
    )
    assert import_response.status_code == 201, import_response.text
    import_payload = import_response.json()
    assert import_payload["status"] == "COMPLETED"
    assert import_payload["summary_json"]["valid"] == 1

    review_response = client.get(f"/api/imports/{import_payload['id']}/review", headers=admin_headers)
    assert review_response.status_code == 200
    review_payload = review_response.json()
    assert len(review_payload) == 1
    assert review_payload[0]["action"] == "ACCEPTED"
    assert review_payload[0]["timetable_entry_id"]

    quality_response = client.get(
        f"/api/imports/{import_payload['id']}/quality-report", headers=admin_headers
    )
    assert quality_response.status_code == 200
    assert quality_response.json()["parser"] == "CsvTimetableParser"
    assert quality_response.json()["confidence"]["HIGH"] == 1

    trace_response = client.get(
        f"/api/imports/{import_payload['id']}/trace", headers=admin_headers
    )
    assert trace_response.status_code == 200
    trace = trace_response.json()["records"][0]
    assert trace["raw_entry_text"] == "English"
    assert trace["interpreted"]["room"] == "414"
    assert trace["source"]["method"] == "CSV_STRUCTURED"

    structures_response = client.get("/api/timetable-structures", headers=admin_headers)
    assert structures_response.status_code == 200
    structure = structures_response.json()[0]
    assert structure["columns"] == ["Monday"]
    assert [(slot["start_time"], slot["end_time"]) for slot in structure["time_slots"]] == [("07:15", "08:15")]
    assert structure["attribute_keys"] == ["batch"]
    assert structure["entries"][0]["attributes_json"] == {"batch": "B1"}

    cached_response = client.post(
        "/api/imports",
        headers=admin_headers,
        files={"file": ("fyba-again.csv", csv_content, "text/csv")},
    )
    assert cached_response.status_code == 201
    assert cached_response.json()["id"] == import_payload["id"]

    slots_response = client.get("/api/time-slots", headers=admin_headers)
    assert slots_response.status_code == 200
    assert [(slot["start_time"], slot["end_time"]) for slot in slots_response.json()] == [("07:15", "08:15")]

    entries_response = client.get("/api/timetable-entries", headers=admin_headers)
    assert entries_response.status_code == 200
    assert entries_response.json()[0]["time_slot_id"] == slots_response.json()[0]["id"]

    schedule_response = client.get(
        "/api/schedule/by-room?room=414&day=Monday&time=07:30",
        headers=admin_headers,
    )
    assert schedule_response.status_code == 200, schedule_response.text
    schedule_payload = schedule_response.json()
    assert schedule_payload[0]["scheduled_status"] == "SCHEDULED"
    assert schedule_payload[0]["subject"] == "English"
    assert schedule_payload[0]["class_name"] == "FYBA"
    assert schedule_payload[0]["division"] == "A"

    availability_response = client.get(
        "/api/availability?room=414&day=Monday&start_time=08:15&end_time=09:15",
        headers=admin_headers,
    )
    assert availability_response.status_code == 200
    assert availability_response.json()["scheduled_availability"] == "AVAILABLE"


def test_import_creates_explicit_classroom_and_maps_lecture(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    csv_content = (
        "class_name,division,day,start_time,end_time,subject,room\n"
        "FYBA,A,Monday,07:15,08:15,Hindi,CC 4\n"
    ).encode()
    response = client.post(
        "/api/imports",
        headers=admin_headers,
        files={"file": ("fyba.csv", csv_content, "text/csv")},
    )
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "COMPLETED"
    assert response.json()["summary_json"]["created_classrooms"] == 1

    classrooms = client.get("/api/classrooms", headers=admin_headers)
    assert classrooms.status_code == 200
    assert [room["room_number"] for room in classrooms.json()] == ["CC 4"]

    reviews = client.get(f"/api/imports/{response.json()['id']}/review", headers=admin_headers).json()
    assert reviews[0]["mapping_status"] == "MATCHED"
    assert reviews[0]["validation_status"] == "VALID"


def test_department_location_creates_department_owned_lab(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    csv_content = (
        "class_name,department,day,start_time,end_time,subject\n"
        "SYBA,Psychology,Monday,09:15,10:15,Psychology (In Department)\n"
    ).encode()
    response = client.post(
        "/api/imports",
        headers=admin_headers,
        files={"file": ("psychology.csv", csv_content, "text/csv")},
    )
    assert response.status_code == 201, response.text
    assert response.json()["status"] == "COMPLETED"

    classroom = client.get("/api/classrooms", headers=admin_headers).json()[0]
    department = client.get("/api/departments", headers=admin_headers).json()[0]
    room_type = client.get("/api/room-types", headers=admin_headers).json()[0]
    assert classroom["room_number"] == "PSYCHOLOGY DEPARTMENT LAB"
    assert classroom["display_name"] == "Psychology Department Lab"
    assert classroom["department_id"] == department["id"]
    assert department["name"] == "Psychology"
    assert room_type["name"] == "Laboratory"


def test_detected_empty_time_slot_is_retained() -> None:
    document = ParsedDocument(
        source_path=Path("sample.pdf"),
        source_type="PDF",
        cells=[
            ParsedCell(
                raw_cell_text="English 414",
                start_time_text="11:15",
                end_time_text="12:15",
                metadata={"time_slot_order": ["11:15-12:15", "12:15-13:15"]},
            )
        ],
    )

    slots = ImportService._detected_time_slots(document)

    assert [(slot[0].strftime("%H:%M"), slot[1].strftime("%H:%M")) for slot in slots] == [
        ("11:15", "12:15"),
        ("12:15", "13:15"),
    ]


def test_timetable_structure_returns_slot_without_entries(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    slot = client.post(
        "/api/time-slots",
        headers=admin_headers,
        json={"start_time": "12:15", "end_time": "13:15"},
    )
    assert slot.status_code == 201, slot.text
    timetable = client.post(
        "/api/timetables",
        headers=admin_headers,
        json={
            "class_name": "FYBA",
            "division": "A",
            "metadata_json": {
                "column_order": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"],
                "time_slot_order": ["12:15-13:15"],
            },
        },
    )
    assert timetable.status_code == 201, timetable.text

    structures = client.get("/api/timetable-structures", headers=admin_headers)

    assert structures.status_code == 200, structures.text
    assert structures.json()[0]["entries"] == []
    assert structures.json()[0]["time_slots"] == [slot.json()]


def test_admin_can_edit_one_timetable_entry_without_reimporting(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    room = client.post(
        "/api/classrooms",
        headers=admin_headers,
        json={"room_number": "CC 4", "display_name": "Classroom CC 4"},
    )
    assert room.status_code == 201, room.text
    imported = client.post(
        "/api/imports",
        headers=admin_headers,
        files={
            "file": (
                "edit.csv",
                b"class_name,division,day,start_time,end_time,subject,room\nFYBA,A,Monday,07:15,08:15,Hindi,CC 4\n",
                "text/csv",
            )
        },
    )
    assert imported.status_code == 201, imported.text
    entry = client.get("/api/timetable-entries", headers=admin_headers).json()[0]
    later_slot = client.post(
        "/api/time-slots",
        headers=admin_headers,
        json={"start_time": "08:15", "end_time": "09:15"},
    ).json()

    updated = client.patch(
        f"/api/timetable-entries/{entry['id']}",
        headers=admin_headers,
        json={
            "day_of_week": "Friday",
            "time_slot_id": later_slot["id"],
            "subject_text": "P.T.",
            "faculty_text": "Sports Department",
            "room_hint": "CC 4",
        },
    )

    assert updated.status_code == 200, updated.text
    payload = updated.json()
    assert payload["day_of_week"] == "Friday"
    assert payload["start_time"] == "08:15"
    assert payload["subject_text"] == "P.T."
    assert payload["classroom_id"] == room.json()["id"]
    assert payload["extraction_method"] == "MANUAL_EDIT"
    allocation = client.get("/api/allocations", headers=admin_headers).json()[0]
    assert allocation["title"] == "P.T."
    assert allocation["day_of_week"] == "Friday"
    assert allocation["start_time"] == "08:15"

    subject = next(item for item in client.get("/api/subjects", headers=admin_headers).json() if item["name"] == "P.T.")
    renamed = client.patch(
        f"/api/subjects/{subject['id']}",
        headers=admin_headers,
        json={"name": "Physical Training", "code": "PT"},
    )
    assert renamed.status_code == 200, renamed.text
    assert client.get("/api/timetable-entries", headers=admin_headers).json()[0]["subject_text"] == "Physical Training"
    assert client.get("/api/allocations", headers=admin_headers).json()[0]["title"] == "Physical Training"
    deleted_subject = client.delete(f"/api/subjects/{subject['id']}", headers=admin_headers)
    assert deleted_subject.status_code == 204
    assert client.get("/api/timetable-entries", headers=admin_headers).json()[0]["subject_id"] is None

    removed = client.delete(f"/api/timetable-entries/{entry['id']}", headers=admin_headers)
    assert removed.status_code == 204
    assert client.get("/api/timetable-entries", headers=admin_headers).json() == []
    assert client.get("/api/allocations", headers=admin_headers).json() == []
    reviews = client.get(f"/api/imports/{imported.json()['id']}/review", headers=admin_headers).json()
    assert reviews[0]["action"] == "IGNORED"
    assert "extraction noise" in reviews[0]["warning"]
