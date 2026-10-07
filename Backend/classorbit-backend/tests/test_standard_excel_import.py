import hashlib
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = PROJECT_ROOT / "assets" / "ClassOrbit_Timetable_Template.xlsx"
DAY_COLUMNS = {
    "Monday": (2, 3),
    "Tuesday": (4, 5),
    "Wednesday": (6, 7),
    "Thursday": (8, 9),
    "Friday": (10, 11),
    "Saturday": (12, 13),
}


def completed_template(
    entries: list[tuple[int, str, str, str]],
    *,
    metadata: bool = True,
    default_room: str = "209",
    custom_times: dict[int, str] | None = None,
) -> bytes:
    workbook = load_workbook(TEMPLATE_PATH)
    sheet = workbook["Timetable Template"]
    if metadata:
        sheet["C3"] = "Computer Science"
        sheet["I3"] = "2026-27"
        sheet["C4"] = "FYBSc CS"
        sheet["I4"] = "A"
        sheet["C5"] = "I"
    sheet["I5"] = default_room
    for row, value in (custom_times or {}).items():
        sheet.cell(row, 1).value = value
    for row, day, subject, room in entries:
        subject_column, room_column = DAY_COLUMNS[day]
        sheet.cell(row, subject_column).value = subject or None
        sheet.cell(row, room_column).value = room or None
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def upload(client: TestClient, headers: dict[str, str], content: bytes, filename: str = "renamed.xlsx"):
    return client.post(
        "/api/imports/standard",
        headers=headers,
        files={"file": (filename, content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )


def create_room(client: TestClient, headers: dict[str, str], number: str, display_name: str | None = None) -> dict:
    response = client.post(
        "/api/classrooms",
        headers=headers,
        json={"room_number": number, "display_name": display_name or f"Room {number}"},
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_exact_template_publishes_valid_rows_and_allocates(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    content = completed_template(
        [(10, "Monday", "DBMS", "Room 209"), (17, "Tuesday", "Python", "")],
        custom_times={17: "2:15 PM - 3:15 PM"},
    )

    response = upload(client, admin_headers, content)
    assert response.status_code == 201, response.text
    payload = response.json()
    summary = payload["summary_json"]
    assert payload["status"] == "COMPLETED"
    assert summary["parser"] == "ClassOrbitTemplateParser"
    assert summary["template_detected"] == "CLASSORBIT_STANDARD_TIMETABLE"
    assert summary["pending_confirmation"] is False
    assert summary["generated_timetable_entries"] == 2
    assert summary["matched_classrooms"] == 2
    assert summary["created_classrooms"] == 1
    assert summary["empty_slots"] == 46
    assert summary["detected_structure"]["columns"] == list(DAY_COLUMNS)
    assert "14:15-15:15" in summary["detected_structure"]["time_slots"]
    assert len(client.get("/api/timetable-entries", headers=admin_headers).json()) == 2
    classrooms = client.get("/api/classrooms", headers=admin_headers).json()
    assert [room["room_number"] for room in classrooms] == ["209"]
    assert classrooms[0]["department_id"] is not None
    assert classrooms[0]["room_type_id"] is not None

    reviews = client.get(f"/api/imports/{payload['id']}/review", headers=admin_headers).json()
    assert {review["day_of_week"] for review in reviews} == {"Monday", "Tuesday"}
    assert reviews[0]["source_sheet"] == "Timetable Template"
    assert all(review["extraction_method"] == "STANDARD_XLSX_TEMPLATE" for review in reviews)
    assert next(review for review in reviews if review["subject_text"] == "Python")["room_hint"] == "209"

    confirmed = client.post(f"/api/imports/{payload['id']}/confirm", headers=admin_headers)
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["created_entries"] == 0
    assert confirmed.json()["status"] == "COMPLETED"
    assert len(client.get("/api/timetable-entries", headers=admin_headers).json()) == 2
    assert len(client.get("/api/allocations", headers=admin_headers).json()) == 2


def test_blank_cells_create_no_entries_but_keep_all_printed_slots(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    response = upload(client, admin_headers, completed_template([]))
    summary = response.json()["summary_json"]
    assert summary["generated_timetable_entries"] == 0
    assert len(summary["detected_structure"]["time_slots"]) == 7
    assert summary["empty_slots"] == 42


def test_room_without_subject_unknown_room_and_multiple_subjects_require_review(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    response = upload(
        client,
        admin_headers,
        completed_template([(10, "Monday", "", "CS Lab 1"), (11, "Tuesday", "Python / Java", "999")]),
    )
    payload = response.json()
    reviews = client.get(f"/api/imports/{payload['id']}/review", headers=admin_headers).json()
    assert len(reviews) == 2
    assert all(review["validation_status"] == "NEEDS_REVIEW" for review in reviews)
    assert set(payload["summary_json"]["unmatched_rooms"]) == {"999", "CS LAB 1"}
    assert any("without Subject/Activity" in (review["warning"] or "") for review in reviews)
    assert any("slash-separated" in (review["warning"] or "") for review in reviews)
    assert client.get("/api/classrooms", headers=admin_headers).json() == []


def test_valid_explicit_lab_is_added_to_classroom_master(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    response = upload(
        client,
        admin_headers,
        completed_template([(10, "Monday", "Programming Practical", "CS Lab 1")]),
    )
    assert response.json()["summary_json"]["created_classrooms"] == 1
    classroom = client.get("/api/classrooms", headers=admin_headers).json()[0]
    assert classroom["room_number"] == "CS LAB 1"
    assert classroom["display_name"] == "Laboratory CS LAB 1"
    room_types = client.get("/api/room-types", headers=admin_headers).json()
    assert next(item for item in room_types if item["id"] == classroom["room_type_id"])["name"] == "Laboratory"


def test_missing_metadata_and_invalid_time_are_reviewed(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    response = upload(
        client,
        admin_headers,
        completed_template([(17, "Monday", "DBMS", "")], metadata=False, custom_times={17: "not a time"}),
    )
    payload = response.json()
    review = client.get(f"/api/imports/{payload['id']}/review", headers=admin_headers).json()[0]
    assert review["validation_status"] == "NEEDS_REVIEW"
    assert "Missing required metadata" in review["warning"]
    assert "Time slot is missing or invalid" in review["warning"]


def test_overlapping_slots_duplicate_class_and_room_conflicts_are_reported(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    create_room(client, admin_headers, "209")
    response = upload(
        client,
        admin_headers,
        completed_template(
            [(10, "Monday", "DBMS", "209"), (17, "Monday", "DBMS", "209")],
            custom_times={17: "7:15 AM - 8:15 AM"},
        ),
    )
    summary = response.json()["summary_json"]
    assert summary["possible_duplicates"] == 1
    assert summary["conflicts"] == 0
    assert summary["needs_review"] >= 1


def test_all_template_sheets_are_imported_with_optional_division_and_parallel_rows(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    workbook = load_workbook(TEMPLATE_PATH)
    first = workbook["Timetable Template"]
    first.title = "Marathi"
    first["C3"] = "Marathi"
    first["I3"] = "2026-27"
    first["C4"] = "TYBA"
    first["I4"] = None
    first["C5"] = None
    first["B10"] = "Marathi Major"
    first["C10"] = "103"
    first["B11"] = "Marathi Elective"
    first["C11"] = "104"
    first["A11"] = first["A10"].value

    second = workbook.copy_worksheet(first)
    second.title = "English"
    second["C3"] = "English"
    second["B10"] = "English Major"
    second["C10"] = "201"
    second["B11"] = "English Elective"
    second["C11"] = "202"

    output = BytesIO()
    workbook.save(output)
    response = upload(client, admin_headers, output.getvalue(), "tyba-all-departments.xlsx")
    assert response.status_code == 201, response.text
    payload = response.json()
    summary = payload["summary_json"]

    assert payload["page_count"] == 2
    assert summary["generated_timetable_entries"] == 4
    assert summary["valid"] == 4
    assert summary["needs_review"] == 0
    assert summary["pending_confirmation"] is False
    assert summary["departments"] == ["English", "Marathi"]
    assert len(client.get("/api/timetable-entries", headers=admin_headers).json()) == 4
    assert len(client.get("/api/classrooms", headers=admin_headers).json()) == 4


def test_authoritative_parallel_room_entries_are_published(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    response = upload(
        client,
        admin_headers,
        completed_template(
            [(10, "Monday", "History", "102"), (11, "Monday", "Performing Arts", "102")],
            custom_times={11: "8:15 AM - 9:15 AM"},
        ),
    )
    payload = response.json()
    summary = payload["summary_json"]

    assert payload["status"] == "COMPLETED"
    assert summary["valid"] == 2
    assert summary["needs_review"] == 0
    assert len(client.get("/api/timetable-entries", headers=admin_headers).json()) == 2


def test_ambiguous_normalized_room_mapping_requires_review(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    create_room(client, admin_headers, "209", "Lecture Hall 209")
    create_room(client, admin_headers, "ROOM 209", "Alternate 209")
    response = upload(client, admin_headers, completed_template([(10, "Monday", "DBMS", "Room 209")]))
    review = client.get(f"/api/imports/{response.json()['id']}/review", headers=admin_headers).json()[0]
    assert review["mapping_status"] == "AMBIGUOUS"
    assert review["validation_status"] == "NEEDS_REVIEW"


def test_renamed_template_is_detected_and_duplicate_upload_is_reused(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    content = completed_template([(10, "Monday", "Online Course", "Online")])
    first = upload(client, admin_headers, content, "college-week-1.xlsx")
    second = upload(client, admin_headers, content, "anything.xlsx")
    assert first.json()["summary_json"]["parser"] == "ClassOrbitTemplateParser"
    assert second.json()["id"] == first.json()["id"]


def test_corrupt_and_invalid_layout_return_friendly_failures(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    corrupt = upload(client, admin_headers, b"not an xlsx", "broken.xlsx")
    assert corrupt.status_code == 201
    assert corrupt.json()["status"] == "FAILED"
    assert "corrupted or unreadable" in corrupt.json()["error_message"]

    workbook = Workbook()
    output = BytesIO()
    workbook.save(output)
    invalid = upload(client, admin_headers, output.getvalue(), "ordinary.xlsx")
    assert invalid.json()["status"] == "FAILED"
    assert "does not match the official ClassOrbit" in invalid.json()["error_message"]


def test_template_download_returns_exact_supplied_workbook(
    client: TestClient,
    admin_headers: dict[str, str],
) -> None:
    response = client.get("/api/timetables/template/download", headers=admin_headers)
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "ClassOrbit_Timetable_Template.xlsx" in response.headers["content-disposition"]
    assert hashlib.sha256(response.content).hexdigest() == hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest()
