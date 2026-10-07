from pathlib import Path

from app.parsers.excel_parser import ExcelTimetableParser
from app.parsers.factory import get_parser_for_file
from app.parsers.matrix_parser import cells_from_matrix
from app.parsers.word_parser import WordTimetableParser


MATRIX = [
    ["Class: FYBA", "Div: A", ""],
    ["Time", "Monday", "Tuesday"],
    ["7.15 am to 8.15 am", "Hindi CC 4", "-"],
    ["8.15 am to 9.15 am", "-", "P.T."],
    ["12.15 pm to 1.15 pm", "-", "-"],
]


def test_matrix_parser_preserves_empty_slots_and_real_cells() -> None:
    cells = cells_from_matrix(
        MATRIX,
        source_name="Sheet1",
        extraction_method="TEST_MATRIX",
        source_page=1,
    )

    assert [(cell.day, cell.raw_cell_text) for cell in cells] == [
        ("Monday", "Hindi CC 4"),
        ("Tuesday", "P.T."),
    ]
    assert cells[0].class_name == "FYBA"
    assert cells[0].division == "A"
    assert cells[0].metadata["time_slot_order"] == [
        "07:15-08:15",
        "08:15-09:15",
        "12:15-13:15",
    ]


def test_matrix_parser_inherits_sheet_metadata_and_lab_section() -> None:
    matrix = [
        ["CLASSORBIT TIMETABLE TEMPLATE", "", ""],
        ["Department", "Animation", ""],
        ["Class / Program", "S.Y.BSc Animation", ""],
        ["Semester", "III", ""],
        ["Time", "Monday", "Tuesday"],
        ["9:00 AM - 10:00 AM", "Theory", "Theory"],
        ["Practical (Pixel Lab)", "", ""],
        ["Time", "Monday", "Tuesday"],
        ["10:00 AM - 12:00 PM", "Digital Editing", "Graphic Design"],
    ]

    cells = cells_from_matrix(
        matrix,
        source_name="Animation",
        extraction_method="TEST_MATRIX",
        source_page=1,
    )

    assert {cell.class_name for cell in cells} == {"S.Y.BSc Animation"}
    assert {cell.metadata["department_name"] for cell in cells} == {"Animation"}
    assert {cell.source_sheet for cell in cells} == {"Animation"}
    practical = [cell for cell in cells if cell.source_row == 9]
    assert {cell.metadata["room_hint"] for cell in practical} == {"Pixel Lab"}
    assert {cell.metadata["location_type"] for cell in practical} == {"LAB"}


def test_matrix_parser_keeps_explicitly_blank_context_fields_blank() -> None:
    matrix = [
        ["CLASSORBIT TIMETABLE TEMPLATE", "", "", "", ""],
        ["Academic Year", "2026-2027", "", "", ""],
        ["Class / Program", "SYBA", "", "Division", ""],
        ["Semester", "", "", "Common Hall / Default Room", ""],
        ["Time", "Monday", "Tuesday", "Wednesday", "Thursday"],
        ["9:00 AM - 10:00 AM", "Marathi", "Hindi", "English", "Economics"],
    ]

    cells = cells_from_matrix(
        matrix,
        source_name="SYBA Marathi (Non-Grant)",
        extraction_method="TEST_MATRIX",
        source_page=1,
    )

    assert {cell.class_name for cell in cells} == {"SYBA"}
    assert {cell.division for cell in cells} == {None}
    assert {cell.metadata.get("semester") for cell in cells} == {None}
    assert {cell.metadata.get("room_hint") for cell in cells} == {None}


def test_excel_visual_grid_is_parsed(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "timetable.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for row in MATRIX:
        sheet.append(row)
    workbook.save(path)

    parsed = ExcelTimetableParser().parse(path)

    assert [(cell.day, cell.raw_cell_text) for cell in parsed.cells] == [
        ("Monday", "Hindi CC 4"),
        ("Tuesday", "P.T."),
    ]
    assert parsed.cells[0].extraction_method == "EXCEL_GRID"


def test_docx_table_is_parsed(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "timetable.docx"
    document = Document()
    table = document.add_table(rows=len(MATRIX), cols=len(MATRIX[0]))
    for row_index, row in enumerate(MATRIX):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value
    document.save(path)

    parsed = WordTimetableParser().parse(path)

    assert [(cell.day, cell.raw_cell_text) for cell in parsed.cells] == [
        ("Monday", "Hindi CC 4"),
        ("Tuesday", "P.T."),
    ]
    assert parsed.cells[0].extraction_method == "DOCX_TABLE"


def test_docx_multiple_tables_keep_nearby_class_context(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "multiple.docx"
    document = Document()
    for division in ("A", "B"):
        document.add_paragraph(f"Class: FYBA Div: {division}")
        table = document.add_table(rows=len(MATRIX) - 1, cols=len(MATRIX[0]))
        for row_index, row in enumerate(MATRIX[1:]):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = value
    document.save(path)

    parsed = WordTimetableParser().parse(path)

    assert sorted({cell.division for cell in parsed.cells}) == ["A", "B"]


def test_factory_accepts_image_and_word_extensions() -> None:
    assert get_parser_for_file(Path("timetable.png")).supported_extensions
    assert get_parser_for_file(Path("timetable.docx")).supported_extensions


def test_macro_enabled_excel_grid_is_supported(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "timetable.xlsm"
    workbook = Workbook()
    sheet = workbook.active
    for row in MATRIX:
        sheet.append(row)
    workbook.save(path)

    parsed = get_parser_for_file(path).parse(path)

    assert parsed.source_type == "EXCEL"
    assert {(cell.class_name, cell.division) for cell in parsed.cells} == {("FYBA", "A")}
    assert len(parsed.cells) == 2
