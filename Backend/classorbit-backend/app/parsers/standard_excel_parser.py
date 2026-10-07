import re
from datetime import date, datetime, time
from pathlib import Path
from zipfile import BadZipFile

from app.parsers.base import ParsedCell, ParsedDocument, ParserCapabilityError
from app.utils.text_normalization import normalize_space
from app.utils.time_utils import normalize_day, parse_time_range, time_to_string


TEMPLATE_NAME = "CLASSORBIT_STANDARD_TIMETABLE"
TEMPLATE_MARKER = "CLASSORBIT TIMETABLE TEMPLATE"
DAY_NAMES = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday")
METADATA_LABELS = {
    "department": "Department",
    "academic_year": "Academic Year",
    "class_program": "Class / Program",
    "division": "Division",
    "semester": "Semester",
    "default_room": "Common Hall / Default Room",
}
REQUIRED_METADATA = ("department", "academic_year", "class_program")


class ClassOrbitTemplateParser:
    supported_extensions = (".xlsx", ".xlsm")

    @classmethod
    def matches(cls, source_path: Path) -> bool:
        workbook = cls._load(source_path)
        return any(cls._layout(sheet) is not None for sheet in workbook.worksheets)

    def parse(self, source_path: Path) -> ParsedDocument:
        workbook = self._load(source_path)
        matches = [
            (sheet, layout)
            for sheet in workbook.worksheets
            if (layout := self._layout(sheet)) is not None
        ]
        if not matches:
            raise ParserCapabilityError(
                "Invalid ClassOrbit timetable layout. Download and use the official ClassOrbit timetable template."
            )

        cells: list[ParsedCell] = []
        warnings: list[str] = []
        slot_order: list[str] = []
        contexts: list[dict] = []
        traces: list[dict] = []
        for sheet, layout in matches:
            sheet_cells, sheet_context, sheet_slots, sheet_warnings = self._parse_sheet(
                workbook, sheet, layout
            )
            cells.extend(sheet_cells)
            contexts.append(sheet_context)
            warnings.extend(sheet_warnings)
            traces.append({
                "sheet": sheet.title,
                "header_row": layout["header_row"],
                "subheader_row": layout["subheader_row"],
            })
            for slot in sheet_slots:
                if slot not in slot_order:
                    slot_order.append(slot)

        first_context = contexts[0]
        return ParsedDocument(
            source_path=source_path,
            source_type="STANDARD_EXCEL",
            cells=cells,
            page_count=len(matches),
            metadata={
                "standard_template": True,
                "template_name": TEMPLATE_NAME,
                "source_sheets": [sheet.title for sheet, _ in matches],
                "department_name": first_context.get("department_name"),
                "academic_year": first_context.get("academic_year"),
                "class_name": first_context.get("class_name"),
                "program_name": first_context.get("program_name"),
                "division": first_context.get("division"),
                "semester": first_context.get("semester"),
                "contexts": contexts,
                "column_order": list(matches[0][1]["day_columns"]),
                "time_slot_order": slot_order,
                "layout": "CLASSORBIT_WEEKLY_SUBJECT_ROOM_V1",
            },
            warnings=warnings,
            parser_name=self.__class__.__name__,
            template_name=TEMPLATE_NAME,
            pages_extracted=len(matches),
            trace={"strategy": TEMPLATE_NAME, "sheets": traces},
        )

    def _parse_sheet(self, workbook, sheet, layout: dict) -> tuple[list[ParsedCell], dict, list[str], list[str]]:
        metadata = {key: self._metadata_value(sheet, label) for key, label in METADATA_LABELS.items()}
        missing_metadata = [METADATA_LABELS[key] for key in REQUIRED_METADATA if not metadata.get(key)]
        slot_order: list[str] = []
        invalid_time_rows: list[int] = []
        rows: list[dict] = []
        slot_occurrences: dict[str, int] = {}

        for row_number in range(layout["subheader_row"] + 1, sheet.max_row + 1):
            raw_time = _display_value(sheet.cell(row_number, layout["time_column"]).value)
            day_values = [
                (
                    day,
                    _display_value(sheet.cell(row_number, columns[0]).value),
                    _display_value(sheet.cell(row_number, columns[1]).value),
                    columns,
                )
                for day, columns in layout["day_columns"].items()
            ]
            has_activity_data = any(subject or room for _, subject, room, _ in day_values)
            if not raw_time and not has_activity_data:
                continue

            start_time, end_time = parse_time_range(raw_time)
            valid_time = bool(start_time and end_time and end_time > start_time)
            if valid_time:
                slot = f"{time_to_string(start_time)}-{time_to_string(end_time)}"
                if slot not in slot_order:
                    slot_order.append(slot)
                if has_activity_data:
                    slot_occurrences[slot] = slot_occurrences.get(slot, 0) + 1
            elif has_activity_data:
                invalid_time_rows.append(row_number)

            rows.append({
                "row_number": row_number,
                "raw_time": raw_time,
                "day_values": day_values,
                "start_time": start_time,
                "end_time": end_time,
                "valid_time": valid_time,
                "slot": slot if valid_time else None,
            })

        cells: list[ParsedCell] = []
        for row in rows:
            row_number = row["row_number"]
            parallel_activity = bool(row["slot"] and slot_occurrences.get(row["slot"], 0) > 1)
            for day, raw_subject, raw_room, (subject_column, room_column) in row["day_values"]:
                if not raw_subject and not raw_room:
                    continue
                room_hint = raw_room or (metadata.get("default_room") if raw_subject else None)
                subject_cell = sheet.cell(row_number, subject_column).coordinate
                room_cell = sheet.cell(row_number, room_column).coordinate
                time_cell = sheet.cell(row_number, layout["time_column"]).coordinate
                warning = None
                if raw_room and not raw_subject:
                    warning = "Classroom/Lab provided without Subject/Activity."
                elif not row["valid_time"]:
                    warning = "Time slot is missing or invalid."
                attributes = {}
                if parallel_activity:
                    attributes = {
                        "parallel_activity": True,
                        "schedule_track": f"{sheet.title} row {row_number}",
                    }
                cells.append(
                    ParsedCell(
                        raw_cell_text=raw_subject,
                        class_name=metadata.get("class_program") or None,
                        division=metadata.get("division") or None,
                        day=day,
                        time_text=row["raw_time"] or None,
                        start_time_text=time_to_string(row["start_time"]),
                        end_time_text=time_to_string(row["end_time"]),
                        source_page=workbook.sheetnames.index(sheet.title) + 1,
                        source_sheet=sheet.title,
                        source_row=row_number,
                        source_column=day,
                        source_cell_reference=f"{sheet.title}!{subject_cell}",
                        extraction_method="STANDARD_XLSX_TEMPLATE",
                        parser_name=self.__class__.__name__,
                        template_name=TEMPLATE_NAME,
                        trace=[{
                            "decision": "standard_template_cell",
                            "sheet": sheet.title,
                            "time_cell": time_cell,
                            "subject_cell": subject_cell,
                            "room_cell": room_cell,
                            "default_room_used": bool(raw_subject and not raw_room and room_hint),
                        }],
                        metadata={
                            "standard_template": True,
                            "academic_year": metadata.get("academic_year") or None,
                            "department_name": metadata.get("department") or None,
                            "program_name": metadata.get("class_program") or None,
                            "semester": metadata.get("semester") or None,
                            "room_hint": room_hint,
                            "location_type": _location_type(room_hint),
                            "required_metadata_missing": missing_metadata,
                            "standard_warning": warning,
                            "attributes": attributes,
                            "raw_subject": raw_subject,
                            "raw_room": raw_room,
                            "time_cell": time_cell,
                            "subject_cell": subject_cell,
                            "room_cell": room_cell,
                            "column_order": list(layout["day_columns"]),
                            "time_slot_order": slot_order,
                            "layout": "CLASSORBIT_WEEKLY_SUBJECT_ROOM_V1",
                            "hall_name": metadata.get("default_room") or None,
                        },
                    )
                )

        warnings: list[str] = []
        if missing_metadata:
            warnings.append(f"{sheet.title}: Missing required metadata: {', '.join(missing_metadata)}.")
        if invalid_time_rows:
            warnings.append(
                f"{sheet.title}: Invalid time values in timetable rows: "
                + ", ".join(str(row) for row in invalid_time_rows)
                + "."
            )
        context = {
            "source_sheet": sheet.title,
            "department_name": metadata.get("department") or None,
            "academic_year": metadata.get("academic_year") or None,
            "class_name": metadata.get("class_program") or None,
            "program_name": metadata.get("class_program") or None,
            "division": metadata.get("division") or None,
            "semester": metadata.get("semester") or None,
            "hall_name": metadata.get("default_room") or None,
            "required_metadata_missing": missing_metadata,
            "time_slot_order": slot_order,
        }
        return cells, context, slot_order, warnings

    @staticmethod
    def _load(source_path: Path):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("ClassOrbit template import requires openpyxl.") from exc
        try:
            return load_workbook(source_path, read_only=False, data_only=True)
        except (BadZipFile, OSError, ValueError, KeyError) as exc:
            raise ParserCapabilityError(
                "The uploaded XLSX workbook is corrupted or unreadable. Download a fresh ClassOrbit template."
            ) from exc

    @classmethod
    def _layout(cls, sheet) -> dict | None:
        marker_found = any(
            _key(cell.value) == _key(TEMPLATE_MARKER)
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 10))
            for cell in row
        )
        for row_number in range(1, min(sheet.max_row, 30) + 1):
            row = list(sheet.iter_rows(min_row=row_number, max_row=row_number))[0]
            time_cell = next((cell for cell in row if _key(cell.value) == "timeslot"), None)
            day_cells = {
                day: next((cell for cell in row if normalize_day(_display_value(cell.value)) == day), None)
                for day in DAY_NAMES
            }
            if not time_cell or any(cell is None for cell in day_cells.values()):
                continue
            subheader_row = row_number + 1
            day_columns: dict[str, tuple[int, int]] = {}
            valid_subheaders = True
            for day, cell in day_cells.items():
                subject_column = cell.column
                room_column = subject_column + 1
                if (
                    _key(sheet.cell(subheader_row, subject_column).value) != "lecturesubject"
                    or _key(sheet.cell(subheader_row, room_column).value) != "classroomlab"
                ):
                    valid_subheaders = False
                    break
                day_columns[day] = (subject_column, room_column)
            if valid_subheaders and marker_found:
                return {
                    "header_row": row_number,
                    "subheader_row": subheader_row,
                    "time_column": time_cell.column,
                    "day_columns": day_columns,
                }
        return None

    @staticmethod
    def _metadata_value(sheet, label: str) -> str:
        target = _key(label)
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 15)):
            for cell in row:
                if _key(cell.value) != target:
                    continue
                end_column = cell.column
                for merged_range in sheet.merged_cells.ranges:
                    if cell.coordinate in merged_range:
                        end_column = merged_range.max_col
                        break
                for column in range(end_column + 1, sheet.max_column + 1):
                    candidate = sheet.cell(cell.row, column).value
                    if _key(candidate) in {_key(item) for item in METADATA_LABELS.values()}:
                        return ""
                    value = _display_value(candidate)
                    if value:
                        return value
                return ""
        return ""


StandardExcelTimetableParser = ClassOrbitTemplateParser


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", normalize_space(str(value or "")).lower())


def _display_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_space(str(value))


def _location_type(room: str | None) -> str | None:
    normalized = _key(room)
    if not normalized:
        return None
    if normalized in {"online", "onlinecourse"}:
        return "ONLINE"
    if normalized in {"indepartment", "department"}:
        return "DEPARTMENT"
    if "lab" in normalized:
        return "LAB"
    return "CLASSROOM"
